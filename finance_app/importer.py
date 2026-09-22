from __future__ import annotations

import hashlib
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Any


class SpreadsheetImportError(ValueError):
    pass


def _date_to_iso(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    raise SpreadsheetImportError(f"Data de emissão inválida: {text!r}")


def _amount_to_cents(value: Any) -> int:
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError) as exc:
        raise SpreadsheetImportError(f"Valor financeiro inválido: {value!r}") from exc
    cents = int(amount * 100)
    if cents <= 0:
        raise SpreadsheetImportError("O valor da nota deve ser maior que zero.")
    return cents


def build_nfse_external_key(transaction_date: str, cnpj: str, amount_cents: int, status: str) -> str:
    raw = f"nfse|{transaction_date}|{cnpj}|{amount_cents}|{status}".encode("utf-8")
    return f"nfse:{hashlib.sha256(raw).hexdigest()[:24]}"


def parse_nfse_workbook(content: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SpreadsheetImportError(
            "A importação Excel precisa do pacote openpyxl. Execute: pip install -r requirements.txt"
        ) from exc

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise SpreadsheetImportError("Não foi possível abrir o arquivo Excel enviado.") from exc

    expected_sheet = "Notas Fiscais - Consolidado"
    if expected_sheet not in workbook.sheetnames:
        raise SpreadsheetImportError(
            f"A aba obrigatória '{expected_sheet}' não foi encontrada."
        )

    sheet = workbook[expected_sheet]
    rows: list[dict[str, Any]] = []

    for row in sheet.iter_rows(min_row=1, values_only=True):
        values = list(row)
        if len(values) < 8:
            continue

        generation = values[1]
        cnpj = str(values[2] or "").strip()
        counterparty = str(values[3] or "").strip()
        municipality = str(values[5] or "").strip()
        amount = values[6]
        raw_status = str(values[7] or "").strip().upper()

        if raw_status not in {"ATIVA", "CANCELADA"} or not cnpj or amount in (None, ""):
            continue

        transaction_date = _date_to_iso(generation)
        amount_cents = _amount_to_cents(amount)
        status = "paid" if raw_status == "ATIVA" else "cancelled"
        rows.append(
            {
                "kind": "income",
                "description": f"NFS-e — {counterparty}",
                "amount_cents": amount_cents,
                "transaction_date": transaction_date,
                "due_date": None,
                "status": status,
                "category": "Honorários e serviços",
                "cost_center": "Jurídico",
                "counterparty": counterparty,
                "counterparty_tax_id": re.sub(r"\D", "", cnpj),
                "document_number": None,
                "notes": f"CNPJ: {cnpj} | Município emissor: {municipality}",
                "source": "import:finance-report",
                "external_key": build_nfse_external_key(transaction_date, cnpj, amount_cents, status),
            }
        )

    if not rows:
        raise SpreadsheetImportError("Nenhuma NFS-e ativa ou cancelada foi encontrada.")
    return rows
